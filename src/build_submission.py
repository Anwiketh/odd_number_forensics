"""Assemble the submission document: EXECUTIVE_SUMMARY + paper + AUDIT, with
figures embedded at their citation points, as a .docx that Google Drive opens
directly as a Google Doc.

Markdown handled: ATX headings, pipe tables, bullet/numbered lists,
blockquotes, fenced code, inline bold / italic / code spans, and a best-effort
LaTeX to Unicode pass (the paper uses \\delta, \\beta, \\rho, \\kappa and a few
display equations).

Figures are inserted immediately after the paragraph that first cites
"Figure N", so the reader meets each graph where it is discussed.
"""
import os, re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PAPER, FIGS = os.path.join(ROOT, "paper"), os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "submission")
os.makedirs(OUT, exist_ok=True)

FIGMAP = {
    1: ("fig1_arm_asymmetry.png",
        "Figure 1 - Per-cell arm asymmetry. The same model and the same incentive "
        "give different disobedience rates depending on which arm you run."),
    2: ("fig2_decomposition.png",
        "Figure 2 - The polarity decomposition: polarity-invariant delta against "
        "content bias beta, with bootstrap CIs over surface variants."),
    3: ("fig3_comprehension.png",
        "Figure 3 - Prior-free comprehension probe. The behavioural ordering "
        "(code, prose, instruction) matches the comprehension ordering exactly."),
    5: ("fig5_geometry.png",
        "Figure 5 - Direction geometry against chance and the split-half "
        "reliability ceiling. The conflict direction transfers across "
        "environments; the content directions sit at chance."),
    6: ("fig6_transfer.png",
        "Figure 6 - Leave-one-environment-out steering at layer 16. Top: the "
        "steering response. Bottom: answer-set mass under the same perturbation "
        "- where it collapses, the intervention has broken the model rather than "
        "changed its mind."),
    8: ("fig8_copying.png",
        "Figure 8 - exp20. Every bar is the same incentive, written with "
        "different numerals."),
    9: ("fig9_payout.png",
        "Figure 9 - exp21. Payout direction x zero-ness x mention order. Every "
        "condition names the same two options."),
    10: ("fig10_scaling.png",
         "Figure 10 - The three effects across six models. The quantity the "
         "environment claims to measure is the smallest of them at every size."),
}

SYM = {"\\delta": "\u03b4", "\\beta": "\u03b2", "\\rho": "\u03c1",
       "\\kappa": "\u03ba", "\\alpha": "\u03b1", "\\times": "\u00d7",
       "\\approx": "\u2248", "\\le": "\u2264", "\\ge": "\u2265",
       "\\ll": "\u226a", "\\sqrt": "\u221a", "\\pm": "\u00b1",
       "\\cdot": "\u00b7", "\\mathbb{E}": "E", "\\tfrac12": "1/2",
       "\\log": "log", "\\sum": "\u03a3", "\\mathrm": "", "\\text": "",
       "\\left": "", "\\right": "", "\\!": "", "\\,": " ", "\\;": " "}


def delatex(t):
    """LaTeX to readable Unicode.

    Order matters. The text-mode wrappers must be unwrapped BEFORE \\frac, or
    the brace pattern stops matching and a ratio silently renders as a product:
    \\log \\frac{P(\\text{disobey})}{P(\\text{obey})} came out as
    "log P(disobey)P(obey)", which states the opposite of the metric.
    """
    t = re.sub(r"\$\$(.+?)\$\$", r"\1", t, flags=re.S)
    t = re.sub(r"\$(.+?)\$", r"\1", t, flags=re.S)
    for cmd in ("text", "mathrm", "mathbb", "mathbf", "operatorname"):
        t = re.sub(r"\\" + cmd + r"\{([^{}]*)\}", r"\1", t)
    for _ in range(3):                      # innermost fraction first
        t2 = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1) / (\2)", t)
        if t2 == t:
            break
        t = t2
    for k, v in SYM.items():
        t = t.replace(k, v)
    t = re.sub(r"_\{?([A-Za-z0-9]+)\}?", r"_\1", t)
    t = re.sub(r"\^\{?([A-Za-z0-9]+)\}?", r"^\1", t)
    t = re.sub(r"\\[A-Za-z]+", "", t)
    return t.replace("{", "").replace("}", "").strip()


CTRL = re.compile("[" + "".join(chr(c) for c in list(range(0,9))+[11,12]+list(range(14,32))) + "]")


def clean(t):
    """docx/lxml rejects control characters outright."""
    return CTRL.sub("", t).replace(chr(9), " ")


def add_runs(par, text):
    text = clean(text)
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+?`)", text)
    for piece in parts:
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            sub = re.split(r"(\*[^*]+?\*)", piece)
            for s in sub:
                if s.startswith("*") and s.endswith("*") and len(s) > 2:
                    par.add_run(s[1:-1]).italic = True
                elif s:
                    par.add_run(s)


def add_figure(doc, n):
    name, cap = FIGMAP[n]
    path = os.path.join(FIGS, name)
    if not os.path.exists(path):
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.4))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(cap)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    return True


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells
             if not all(re.fullmatch(r":?-{2,}:?", (x or "-")) for x in c)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    try:
        t.style = "Table Grid"
    except KeyError:
        pass
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        rc = t.add_row().cells
        for j in range(ncol):
            txt = delatex(row[j]) if j < len(row) else ""
            add_runs(rc[j].paragraphs[0], txt)
            for pr in rc[j].paragraphs:
                for run in pr.runs:
                    run.font.size = Pt(8.5)
                    if i == 0:
                        run.bold = True
    doc.add_paragraph()


def render(doc, md, shown, base_level=1):
    lines = md.split("\n")
    i = 0
    buf = []

    def flush():
        if not buf:
            return
        text = delatex(" ".join(buf).strip())
        del buf[:]
        if not text:
            return
        p = doc.add_paragraph()
        add_runs(p, text)
        for n in sorted(FIGMAP):
            if n not in shown and re.search(r"Figure " + str(n) + r"\b", text):
                if add_figure(doc, n):
                    shown.add(n)

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("|"):
            flush()
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i])
                i += 1
            add_table(doc, tbl)
            continue
        if s.startswith("```"):
            flush()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(code))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", ln)
        if m:
            flush()
            lvl = min(len(m.group(1)) + base_level, 5)
            h = delatex(m.group(2))
            h = re.sub(r"\*+([^*]+)\*+", r"", h).replace("`", "")
            doc.add_heading(h, level=lvl)
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", ln):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, delatex(re.sub(r"^\s*[-*]\s+", "", ln)))
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", ln):
            flush()
            p = doc.add_paragraph(style="List Number")
            add_runs(p, delatex(re.sub(r"^\s*\d+\.\s+", "", ln)))
            i += 1
            continue
        if s.startswith(">"):
            flush()
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, delatex(s.lstrip("> ")))
            i += 1
            continue
        if s in ("---", "***", "___") or not s:
            flush()
            i += 1
            continue
        buf.append(s)
        i += 1
    flush()


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)

    shown = set()

    doc.add_heading("What Is the Odd Number Environment Actually Measuring?", 0)
    sub = doc.add_paragraph()
    r = sub.add_run("A forensic audit of a reward-hacking probe   |   "
                    "MATS 12.0 application, Neel Nanda stream")
    r.italic = True
    r.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    doc.add_heading("Executive summary", level=1)
    ex = open(os.path.join(PAPER, "EXECUTIVE_SUMMARY.md"), encoding="utf-8").read()
    ex = re.sub(r"^#\s+.*\n", "", ex, count=1)
    render(doc, ex, shown, base_level=0)
    if 10 not in shown and add_figure(doc, 10):
        shown.add(10)

    # He asks for randomly selected raw examples immediately after the summary.
    ex_path = os.path.join(OUT, "random_examples.md")
    if os.path.exists(ex_path):
        doc.add_page_break()
        doc.add_heading("Randomly sampled raw examples", level=1)
        p0 = doc.add_paragraph()
        add_runs(p0, "Everything here rests on the prompts being what I say they "
                     "are and on the read-out looking at the tokens it claims to. "
                     "Neither is visible from a summary statistic, so these are "
                     "drawn uniformly at random from the design, seeded and "
                     "reproducible, and shown with the full answer distribution.")
        raw = open(ex_path, encoding="utf-8").read()
        raw = re.sub("^#" + chr(92) + "s+.*" + chr(92) + "n", "", raw, count=1)
        render(doc, raw, shown, base_level=1)

    doc.add_page_break()
    doc.add_heading("Full write-up", level=1)
    pm = open(os.path.join(PAPER, "paper.md"), encoding="utf-8").read()
    pm = re.sub(r"^#\s+.*\n", "", pm, count=1)
    render(doc, pm, shown, base_level=0)

    leftover = [n for n in sorted(FIGMAP) if n not in shown]
    if leftover:
        doc.add_page_break()
        doc.add_heading("Additional figures", level=1)
        for n in leftover:
            if add_figure(doc, n):
                shown.add(n)

    doc.add_page_break()
    doc.add_heading("Appendix - what I verified, and what I got wrong", level=1)
    p = doc.add_paragraph()
    add_runs(p, "This is the project's self-audit log: six rounds of checking my "
                "own numbers against the raw results, three retractions of my own "
                "headline, and a live defect found in my own experiment code. It "
                "is included because the checking is part of the work.")
    render(doc, open(os.path.join(PAPER, "AUDIT.md"), encoding="utf-8").read(),
           shown, base_level=0)

    out = os.path.join(OUT, "MATS12_OddNumber_Submission.docx")
    doc.save(out)
    print("wrote", out)
    print("figures embedded:", sorted(shown))


if __name__ == "__main__":
    main()
